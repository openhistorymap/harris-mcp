"""XLSX adapter — paired tables in a single workbook.

Read: looks for sheets named (case-insensitively) ``contexts`` and
``relations``. Write: emits exactly those two sheets.
"""

from __future__ import annotations

from pathlib import Path

from ..model import Context, Matrix, Relation
from ..registry import new_id


def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for XLSX support (`pip install openpyxl`)"
        ) from exc


def load(path: str | Path) -> Matrix:
    _require_openpyxl()
    from openpyxl import load_workbook

    p = Path(path)
    wb = load_workbook(p, data_only=True, read_only=True)
    by_name = {s.title.lower(): s for s in wb.worksheets}
    if "contexts" not in by_name or "relations" not in by_name:
        raise ValueError(
            f"{p} must contain sheets named 'contexts' and 'relations'"
        )
    m = Matrix(id=new_id(hint=p.stem), name=p.stem, path=str(p))

    def _rows(ws):
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
        for row in rows:
            yield {header[i]: row[i] for i in range(len(header)) if header[i]}

    known = {"type", "period", "phase", "group", "description"}
    for row in _rows(by_name["contexts"]):
        ctx_id = row.get("id") or row.get("ID")
        if ctx_id is None:
            continue
        ctx_id = str(ctx_id)
        kwargs = {k: row[k] for k in known if row.get(k) not in (None, "")}
        extras = {k: v for k, v in row.items() if k not in known and k.lower() != "id" and v not in (None, "")}
        m.add_context(Context(id=ctx_id, attrs=extras, **kwargs))

    for row in _rows(by_name["relations"]):
        a = row.get("a") or row.get("from") or row.get("source")
        b = row.get("b") or row.get("to") or row.get("target")
        kind = row.get("kind") or row.get("relation") or "above"
        if a is None or b is None:
            continue
        m.add_relation(Relation(a=str(a), b=str(b), kind=str(kind)))

    m.rebuild_indexes()
    return m


def dump(matrix: Matrix, path: str | Path) -> None:
    _require_openpyxl()
    from openpyxl import Workbook

    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ctx_sheet = wb.active
    ctx_sheet.title = "contexts"
    ctx_fields = ["id", "type", "period", "phase", "group", "description"]
    extra_keys: list[str] = []
    for c in matrix.contexts.values():
        for k in c.attrs:
            if k not in extra_keys:
                extra_keys.append(k)
    ctx_sheet.append(ctx_fields + extra_keys)
    for c in matrix.contexts.values():
        row = [getattr(c, k) or "" for k in ctx_fields]
        row.extend([c.attrs.get(k, "") for k in extra_keys])
        ctx_sheet.append(row)

    rel_sheet = wb.create_sheet("relations")
    rel_sheet.append(["a", "b", "kind"])
    for r in matrix.relations:
        rel_sheet.append([r.a, r.b, r.kind])

    wb.save(p)
